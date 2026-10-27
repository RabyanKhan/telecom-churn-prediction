
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


# =====================================================
# 1. PROJECT PATHS AND PRESENTATION SETTINGS
# =====================================================

ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"
PLOTS = REPORTS / "plots"
SCREENSHOTS = REPORTS / "screenshots"

OUT = REPORTS / "Telecom_Churn_Prediction_Presentation.pptx"

# Create reports folder if it does not exist
REPORTS.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Professional theme colors
NAVY = RGBColor(22, 43, 77)
BLUE = RGBColor(39, 100, 176)
TEAL = RGBColor(28, 145, 140)
DARK = RGBColor(45, 52, 60)
LIGHT = RGBColor(242, 246, 250)
WHITE = RGBColor(255, 255, 255)
GREY = RGBColor(105, 115, 125)


# =====================================================
# 2. REUSABLE DESIGN FUNCTIONS
# =====================================================

def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(
        Inches(0.65), Inches(0.35),
        Inches(12), Inches(0.65)
    )

    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos Display"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY

    # Accent line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.68), Inches(1.08),
        Inches(1.25), Inches(0.055)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    if subtitle:
        box = slide.shapes.add_textbox(
            Inches(0.68), Inches(1.2),
            Inches(11.9), Inches(0.35)
        )
        p = box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.name = "Aptos"
        p.font.size = Pt(12)
        p.font.color.rgb = GREY


def footer(slide, number):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0, Inches(7.28),
        prs.slide_width, Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    box = slide.shapes.add_textbox(
        Inches(0.65), Inches(7.02),
        Inches(12), Inches(0.2)
    )
    p = box.text_frame.paragraphs[0]
    p.text = f"Telecom Customer Churn Prediction | {number}"
    p.font.name = "Aptos"
    p.font.size = Pt(9)
    p.font.color.rgb = GREY


def bullets(slide, items, y=1.6, font=19,
            x=0.9, w=11.6, h=5.1):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "\u2022  " + item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(font)
        p.font.color.rgb = DARK
        p.space_after = Pt(15)


def metric(slide, x, value, label,
           color=BLUE, w=2.8):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(1.75),
        Inches(w), Inches(1.4)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = color

    box = slide.shapes.add_textbox(
        Inches(x + 0.08), Inches(1.95),
        Inches(w - 0.16), Inches(0.45)
    )
    p = box.text_frame.paragraphs[0]
    p.text = value
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos Display"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = color

    box = slide.shapes.add_textbox(
        Inches(x + 0.08), Inches(2.55),
        Inches(w - 0.16), Inches(0.3)
    )
    p = box.text_frame.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(12)
    p.font.color.rgb = DARK


def chart_slide(heading, filename, insight, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, heading)

    path = PLOTS / filename

    if path.exists():
        slide.shapes.add_picture(
            str(path),
            Inches(0.55), Inches(1.4),
            width=Inches(8.1),
            height=Inches(5.0)
        )
    else:
        bullets(
            slide,
            [f"Chart not found: reports/plots/{filename}"],
            y=2.0
        )

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8.95), Inches(2.0),
        Inches(3.65), Inches(2.8)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = TEAL

    box = slide.shapes.add_textbox(
        Inches(9.2), Inches(2.25),
        Inches(3.15), Inches(2.3)
    )
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "KEY INSIGHT"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = TEAL

    p = tf.add_paragraph()
    p.text = insight
    p.font.name = "Aptos"
    p.font.size = Pt(17)
    p.font.color.rgb = DARK
    p.space_before = Pt(12)

    footer(slide, number)


# =====================================================
# 3. SLIDE 1 — COVER
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)

accent = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE,
    0, 0, Inches(0.18), prs.slide_height
)
accent.fill.solid()
accent.fill.fore_color.rgb = TEAL
accent.line.fill.background()

box = slide.shapes.add_textbox(
    Inches(0.9), Inches(1.6),
    Inches(11.5), Inches(2.1)
)
p = box.text_frame.paragraphs[0]
p.text = "TELECOM CUSTOMER\nCHURN PREDICTION"
p.font.name = "Aptos Display"
p.font.size = Pt(38)
p.font.bold = True
p.font.color.rgb = WHITE

box = slide.shapes.add_textbox(
    Inches(0.95), Inches(4.15),
    Inches(11), Inches(0.8)
)
p = box.text_frame.paragraphs[0]
p.text = (
    "Exploratory Data Analysis | Machine Learning | "
    "Interactive Prediction"
)
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(204, 224, 238)

box = slide.shapes.add_textbox(
    Inches(0.95), Inches(6.4),
    Inches(10), Inches(0.3)
)
p = box.text_frame.paragraphs[0]
p.text = "PROJECT PRESENTATION"
p.font.size = Pt(13)
p.font.color.rgb = RGBColor(204, 224, 238)


# =====================================================
# 4. SLIDE 2 — PROJECT OVERVIEW
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Project Overview", "Business problem and project goals")

bullets(slide, [
    "Business problem: customer churn can reduce recurring revenue and increase acquisition costs.",
    "Goal: explore customer behavior and predict which customers may leave the telecom service.",
    "Workflow: data audit and cleaning, EDA, feature engineering, model training, tuning and evaluation.",
    "Deliverable: an interactive prediction prototype to support customer-retention analysis."
])

footer(slide, 2)


# =====================================================
# 5. SLIDE 3 — DATASET AND DATA QUALITY
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Dataset & Data Quality", "Data quality checks before model training")

bullets(slide, [
    "Raw dataset inspected: 1,300 rows and 23 columns.",
    "Audit identified duplicate records, inconsistent category labels, placeholder values, mixed numeric formats and missing data.",
    "Cleaning standardized text and churn labels, converted numeric fields and checked invalid values.",
    "Duplicate rows were removed; the original raw dataset was preserved.",
    "Cleaned dataset: telecom_churn_clean.csv.",
    "Modeling dataset: 1,238 rows after cleaning and target-availability checks."
], font=17)

footer(slide, 3)


# =====================================================
# 6. SLIDES 4–8 — EXPLORATORY DATA ANALYSIS
# =====================================================

chart_slide(
    "EDA: Churn Distribution",
    "churn_distribution.png",
    "403 of 1,238 customers are labeled churned (32.6%), making churn an important business concern.",
    4
)

chart_slide(
    "EDA: Churn by Contract",
    "churn_by_contract.png",
    "Month-to-month customers have the highest observed churn rate (50.7%) among contract types.",
    5
)

chart_slide(
    "EDA: Monthly Charges",
    "monthly_charges_vs_churn.png",
    "Median monthly charges are higher for churned customers (84.80) than retained customers (71.27).",
    6
)

chart_slide(
    "EDA: Tenure and Churn",
    "tenure_vs_churn.png",
    "Churned customers have shorter median tenure (18 months) than retained customers (27.5 months).",
    7
)

chart_slide(
    "EDA: Churn by Internet Service",
    "churn_by_internet_service.png",
    "Fiber optic customers have the highest observed churn rate (41.4%) among internet-service groups.",
    8
)


# =====================================================
# 7. SLIDE 9 — FEATURE ENGINEERING
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Feature Engineering & Preprocessing")

bullets(slide, [
    "Engineered features include customer_risk_signals and calculated_monthly_charge.",
    "Numeric preprocessing: median imputation and standard scaling.",
    "Categorical preprocessing: most-frequent imputation and one-hot encoding.",
    "Preprocessing steps are contained within model pipelines to reduce data leakage.",
    "The train/test split uses stratification to preserve target-class proportions."
], font=18)

footer(slide, 9)


# =====================================================
# 8. SLIDE 10 — MODEL DEVELOPMENT
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Model Development & Evaluation")

bullets(slide, [
    "Baseline classifiers: Logistic Regression, Decision Tree, Random Forest and XGBoost.",
    "Random Forest and XGBoost were tuned using cross-validation.",
    "Evaluation metrics include accuracy, precision, recall, F1 score and ROC-AUC.",
    "Confusion matrices were generated for both tuned models.",
    "Recall and F1 are useful churn metrics because missing likely churners may be costly."
], font=18)

footer(slide, 10)


# =====================================================
# 9. SLIDE 11 — TUNED RANDOM FOREST RESULTS
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(
    slide,
    "Tuned Random Forest: Test Results",
    "Selected for the highest test F1 score among the compared models"
)

metric(slide, 0.7, "73.39%", "Accuracy")
metric(slide, 3.85, "57.89%", "Precision", TEAL)
metric(slide, 7.0, "67.90%", "Recall")
metric(slide, 10.05, "62.50%", "F1 Score", TEAL, w=2.55)

bullets(slide, [
    "ROC-AUC: 79.27%.",
    "Tuned XGBoost achieved 74.60% accuracy, but its F1 score was lower at 55.94%.",
    "These results depend on the dataset and test split; further validation is recommended before real-world use."
], y=3.65, font=16, h=2.6)

footer(slide, 11)


# =====================================================
# 10. SLIDES 12–13 — CONFUSION MATRICES
# =====================================================

confusion_slides = [
    (
        "Tuned Random Forest: Confusion Matrix",
        "random_forest_tuned_confusion_matrix.png",
        12,
        "Selected model. Review false negatives and false positives when planning retention actions."
    ),
    (
        "Tuned XGBoost: Confusion Matrix",
        "xgboost_tuned_confusion_matrix.png",
        13,
        "Comparison model. Its test F1 score is lower than tuned Random Forest in this experiment."
    )
]

for heading, filename, number, note in confusion_slides:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, heading)

    path = PLOTS / filename

    if path.exists():
        slide.shapes.add_picture(
            str(path),
            Inches(0.55), Inches(1.4),
            width=Inches(8.1),
            height=Inches(5.0)
        )
    else:
        bullets(
            slide,
            [f"Chart not found: reports/plots/{filename}"],
            y=2.0
        )

    box = slide.shapes.add_textbox(
        Inches(9.0), Inches(2.0),
        Inches(3.4), Inches(2.8)
    )
    p = box.text_frame.paragraphs[0]
    p.text = note
    p.font.name = "Aptos"
    p.font.size = Pt(17)
    p.font.color.rgb = DARK

    footer(slide, number)


# =====================================================
# 11. SLIDE 14 — STREAMLIT DASHBOARD
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Interactive Prediction Dashboard")

bullets(slide, [
    "Built with Streamlit to accept customer details through an interactive form.",
    "The saved machine-learning pipeline generates a churn prediction and probability estimate.",
    "Screenshot file: reports/screenshots/prediction_result.png.",
    "A prediction is a model estimate, not a guarantee of future behavior."
], y=1.65, font=17, x=0.75, w=5.5, h=4.8)

screenshot = SCREENSHOTS / "prediction_result.png"

if screenshot.exists():
    slide.shapes.add_picture(
        str(screenshot),
        Inches(6.4), Inches(1.65),
        width=Inches(6.2)
    )
else:
    box = slide.shapes.add_textbox(
        Inches(6.5), Inches(2.2),
        Inches(5.8), Inches(1.0)
    )
    p = box.text_frame.paragraphs[0]
    p.text = "Dashboard screenshot not found."
    p.font.size = Pt(18)
    p.font.color.rgb = GREY

footer(slide, 14)


# =====================================================
# 12. SLIDE 15 — CONCLUSION
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Conclusion & Next Steps")

bullets(slide, [
    "Month-to-month contracts, higher monthly charges and shorter tenure are associated with higher observed churn in this dataset.",
    "Tuned Random Forest achieved the highest test F1 score among the evaluated models.",
    "The Streamlit application demonstrates how a trained model can support customer-retention workflows.",
    "Next steps: validate on newer data, monitor performance, review fairness and assess probability calibration before deployment."
], font=17)

footer(slide, 15)


# =====================================================
# 13. SLIDE 16 — PROJECT DELIVERABLES
# =====================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Project Deliverables")

bullets(slide, [
    "telecom_churn_raw.csv — original dataset.",
    "telecom_churn_clean.csv — cleaned dataset.",
    "reports/ — audit, missing-value, model comparison and recommendation outputs.",
    "reports/plots/ — five EDA charts and two tuned-model confusion matrices.",
    "models/ — saved baseline and tuned model pipelines.",
    "app.py — interactive Streamlit prediction dashboard.",
    "README.md and requirements.txt — project documentation and dependencies."
], font=16)

footer(slide, 16)


# =====================================================
# 14. SAVE PRESENTATION
# =====================================================

prs.save(OUT)

print()
print("=" * 60)
print("PRESENTATION CREATED SUCCESSFULLY")
print("=" * 60)
print(f"File: {OUT}")
print(f"Number of slides: {len(prs.slides)}")
print()
print("Open the PowerPoint file and review all slides before submission.")
print("Check that the charts and dashboard screenshot appear correctly.")
